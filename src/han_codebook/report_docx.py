"""Readable Word export of a completed, explicitly unverified research analysis.

No secrets, API prompts or hidden synthetic ground truth are included.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _table(doc: Document, headings: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headings))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(headings):
        cell = table.rows[0].cells[i]
        cell.text = title
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    doc.add_paragraph()


def investigation_docx(result: dict[str, Any], events: list[dict[str, Any]]) -> bytes:
    """Return a DOCX payload; original verified mappings are never replaced."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft JhengHei"
    styles["Normal"].font.size = Pt(9)
    for key in ("Title", "Heading 1", "Heading 2"):
        styles[key].font.name = "Microsoft JhengHei"

    title = doc.add_heading("中文歷史詞組密電｜研究分析報告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("注意：尚未驗證的模型與人工假說，不等於密電破解成果。")
    proposals = result.get("proposals", {})
    verified = sum(p.get("status") == "verified_alignment" for p in proposals.values())
    _table(doc, ["項目", "內容"], [
        ["案例", str(result.get("case_id", ""))],
        ["分析電文", str(len(result.get("decoded_messages", [])))],
        ["已驗證碼組", str(verified)],
        ["待驗證碼組", str(len(proposals) - verified)],
        ["完成 LLM 回合", str(result.get("agent_rounds_completed", 0))],
        ["使用元數據", str(result.get("metadata_used", False))],
    ])
    baseline = next((e.get("payload", {}).get("mapping", {}) for e in events if e.get("stage") == "statistical_baseline"), {})
    scores = next((e.get("payload", {}).get("score", "") for e in events if e.get("stage") == "statistical_baseline"), "")
    round_rows = [["Round 0 (統計基線)", "—", "—", "—", str(scores)]]
    suggested = {x.get("round"): x.get("payload", {}) for x in events if x.get("stage") == "llm_proposals_unverified"}
    for e in events:
        if e.get("stage") != "rescore":
            continue
        payload = e.get("payload", {})
        sent = suggested.get(e.get("round"), {})
        round_rows.append([str(e.get("round", "")), str(len(sent.get("accepted", []))),
                           str(sent.get("rejected_count", 0)), str(payload.get("changed_count", "")),
                           str(payload.get("score", ""))])
    doc.add_heading("一、逐輪分析", level=1)
    _table(doc, ["回合", "採納假說", "拒絕", "首選變更", "內部目標分數"], round_rows)
    doc.add_paragraph("內部目標分數不能視為機率或正確率；不同設定間亦不適合直接比較。")
    doc.add_heading("二、碼組候選與已知明文", level=1)
    _table(doc, ["碼組", "驗證狀態", "統計初選", "最終候選", "是否變動"], [
        [code, "已驗證明文" if p.get("status") == "verified_alignment" else "未驗證假說",
         str(baseline.get(code, "")), str(p.get("term", "")),
         "是" if baseline.get(code) != p.get("term") else "否"]
        for code, p in sorted(proposals.items())
    ])
    overrides = result.get("bold_hypotheses", {})
    if overrides:
        doc.add_heading("三、研究員大膽指定（反事實假說）", level=1)
        doc.add_paragraph("以下為研究員自行提出的反對假說，不會覆寫可信明密文證據，也未作為本次求解器的 verified anchor。")
        _table(doc, ["碼組", "已驗證原詞", "研究員替代假說", "理由"], [
            [code, str(result.get("verified_anchors", {}).get(code, "")), str(item.get("term", "")),
             str(item.get("note", ""))]
            for code, item in sorted(overrides.items())
        ])
    doc.add_heading("四、逐封電文候選解讀", level=1)
    _table(doc, ["電文", "推測領域", "目前候選明文", "已驗證碼組"], [
        [str(m.get("id", "")), str(m.get("domain_inferred") or "未知"), str(m.get("reading", "")),
         f"{m.get('verified_codes',0)}/{m.get('total_codes',0)}"]
        for m in result.get("decoded_messages", [])
    ])
    doc.add_heading("五、研究限制", level=1)
    for line in (
        "僅處理固定分隔數字碼組，不代表可破解所有真實電碼系統。",
        "詞彙先驗、收發者和目擊情資都是不確定的推理條件。",
        "模型候選與研究員大膽指定均非已驗證明文；不能改寫可靠原始紀錄。",
        "應以推理過程完全隔離的獨立標準答案，才能評估真實正確率。",
    ):
        doc.add_paragraph(line, style="List Bullet")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
