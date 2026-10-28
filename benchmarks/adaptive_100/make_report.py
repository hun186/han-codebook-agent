"""Generate publishable Word report, including every one of the 100 ciphertexts.

Usage: python benchmarks/adaptive_100/make_report.py --data-dir benchmarks/adaptive_100 --out vercel-demo/reports/adaptive_100_full_report.docx
The key is scoring-only and is never bundled with a browser inference input.
"""
from __future__ import annotations
import argparse, json, hashlib, collections
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONT='Noto Sans CJK TC'
NAVY=RGBColor(27,52,83)
TEAL=RGBColor(0,108,122)
GREY=RGBColor(93,106,120)
RED=RGBColor(169,45,50)

def p(doc,text,*,size=9,bold=False,color=None,space=3):
    line=doc.add_paragraph()
    line.paragraph_format.space_after=Pt(space)
    r=line.add_run(str(text));r.font.name=FONT;r.font.size=Pt(size);r.bold=bold
    if color:r.font.color.rgb=color
    return line

def table(doc,head,rows,fontsize=8):
    t=doc.add_table(rows=1,cols=len(head));t.style='Table Grid'
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,name in enumerate(head):
        cell=t.rows[0].cells[i];cell.text=str(name)
        for run in cell.paragraphs[0].runs:run.font.name=FONT;run.font.size=Pt(fontsize);run.bold=True
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E3F0F3')
        cell._tc.get_or_add_tcPr().append(sh)
    rowhead=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rowhead)
    for j,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val)
            for par in cells[i].paragraphs:
                par.paragraph_format.space_after=Pt(0)
                for run in par.runs:run.font.name=FONT;run.font.size=Pt(fontsize)
            if j%2:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F5F8FA')
                cells[i]._tc.get_or_add_tcPr().append(sh)
    return t

def render(base:Path,out:Path):
    public=json.loads((base/'han_adaptive_new_public.json').read_text(encoding='utf-8'))
    locked_path=base/'han_adaptive_new_predictions_LOCKED.json'
    lock=json.loads(locked_path.read_text(encoding='utf-8'))
    gold=json.loads((base/'reference'/'han_adaptive_new_gold_PRIVATE.json').read_text(encoding='utf-8'))
    noaux=json.loads((base/'han_adaptive_no_aux_LOCKED.json').read_text(encoding='utf-8'))
    rounds=lock['rounds'];final=rounds[2]['predictions']
    truth=gold['code_to_word'];msgs=public['messages']
    freq=collections.Counter(c for m in msgs for c in m['codes'])
    assert len(msgs)==100 and len(truth)==60 and sum(freq.values())==852
    assert len(rounds)==3 and len(final)==60
    summaries=[]
    for stage in [*rounds,{'stage':'no_aux','predictions':noaux['predictions'],'vocabulary':rounds[-1]['vocabulary']}]:
        mapping=stage['predictions'];right={c for c in truth if mapping[c]==truth[c]}
        summaries.append((stage['stage'],len(stage['vocabulary']),len(right),sum(freq[c] for c in right),
                          sum(all(mapping[c]==truth[c] for c in m['codes']) for m in msgs)))
    assert [(x[2],x[3],x[4]) for x in summaries]==[(8,156,0),(28,444,0),(31,470,0),(4,50,0)]
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=doc.styles[name];style.font.name=FONT;style.font.color.rgb=NAVY
    doc.styles['Normal'].font.size=Pt(9)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].font.color.rgb=TEAL
    doc.styles['Heading 2'].font.size=Pt(11)
    hd=sec.header.paragraphs[0];hd.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    x=hd.add_run('漢字碼本研究｜100 封合成密電　·　未經歷史解密驗證')
    x.font.name=FONT;x.font.size=Pt(8);x.font.color.rgb=GREY
    ft=sec.footer.paragraphs[0];ft.alignment=WD_ALIGN_PARAGRAPH.CENTER
    ft.add_run('Synthetic research only　｜　')
    num=OxmlElement('w:fldSimple');num.set(qn('w:instr'),'PAGE');ft._p.append(num)
    doc.add_heading('逐輪動態詞庫與跨報輔助情報\n100 封合成密電測試報告',0)
    p(doc,'60 個碼組的對應正確率　｜　852 處逐碼對照　｜　100 封逐報附錄',size=11,bold=True,color=TEAL,space=12)
    p(doc,'2026-10-28　研究資料 N-001～N-100　｜　模擬案例、非真實軍事情報',size=9,space=8)
    p(doc,'研究限制：GPT-6 僅協助提出動態候選詞方向；真正的碼組分配是離線 Python 一對一指派程式，未呼叫 GPT API，也沒有進行人工覆核。此合成樣本使用對求解器有利的固定角色位置，不能推論真實密電破譯成功率。',size=9.5,color=RED,space=10)
    doc.add_heading('重點結果',1)
    p(doc,'第三輪不同碼組 Top-1 正確率 31/60（51.67%），逐碼位置 470/852（55.16%），完整還原 0/100 封。固定詞庫與動態詞庫是在同一份合成資料比較；和先前已揭盲的另一個 100 封案例不得直接作 A/B。')
    table(doc,['階段','候選詞','碼組 Top-1','位置 Top-1','整封正確'],[
       ['固定詞庫','27','8/60（13.33%）','156/852（18.31%）','0/100'],
       ['第2輪：情報擴充','48','28/60（46.67%）','444/852（52.11%）','0/100'],
       ['第3輪：語境擴充','77','31/60（51.67%）','470/852（55.16%）','0/100'],
       ['去除輔助訊號','77','4/60（6.67%）','50/852（5.87%）','0/100']],8.5)
    doc.add_heading('資料、方法與解釋限制',1)
    for entry in [
        '資料包含 100 封新合成密電、60 個隨機數字碼組及 34 則觀察；每封有寄件／收件者和日期，觀察僅為假說，不是逐碼已知明文。',
        '固定詞庫 27 詞，第2輪自公開收發單位與觀察自動增加 21 詞，第3輪加入事先提出的 29 個領域詞；不看 gold 的離線求解器分配各碼組。',
        '已知合成角色結構「單位／動作／物件／地點／時間／狀態」；這一資訊在真實歷史材料中未必能得到，且提高了本次正確率。',
        '第三輪候選包含所有 60 個真正詞彙，卻仍有 29 個不同碼組錯誤。詞彙覆蓋不等於碼組對應正確。',
        '消融試驗將輔助訊號移除後，碼組正確率為 4/60（6.67%）；該差距只反映本資料設計的強相關特性。',
        '參考答案已揭曉，本案例不能再次作為獨立盲測；公開展示只能重播鎖定預測，不能宣稱新一輪真實 LLM 推理。'
    ]:p(doc,'• '+entry,size=8.8,space=3)
    p(doc,'鎖定預測 SHA256：'+hashlib.sha256(locked_path.read_bytes()).hexdigest(),size=7.7,color=GREY)
    doc.add_page_break()
    doc.add_heading('附錄 A｜60 個碼組完整對照表',1)
    p(doc,'正確明文僅揭盲後提供給評分和此報告；推論階段不允許讀取此欄。',size=8.4,color=GREY)
    table(doc,['碼組','頻次','固定詞庫','第2輪','第3輪','正確明文','判定'],[
      [c,freq[c],*[r['predictions'][c] for r in rounds],truth[c],
       '✓' if final[c]==truth[c] else '✗'] for c in sorted(truth,key=int)],7.75)
    doc.add_page_break()
    doc.add_heading('附錄 B｜100 封原密電與明文逐碼對照',1)
    p(doc,'每封逐一列出日期、收發者、第三方觀察、原始密碼組、第三輪還原、正確明文，以及依原始碼組序列對位的錯誤標記。每個「✓」為字碼對應完全相同；「✗」為錯詞或未解。',size=8.5,space=8)
    count_correct=count_total=0
    for index,m in enumerate(msgs,1):
        codes=m['codes'];words=gold['messages_words'][m['id']]
        assert len(codes)==len(words)
        guessed=[final[c] for c in codes]
        hit=[guess==word for guess,word in zip(guessed,words)]
        count_correct+=sum(hit);count_total+=len(codes)
        title=p(doc,f"{index:03d}　{m['id']}　｜　正確 {sum(hit)}/{len(codes)}（{100*sum(hit)/len(codes):.1f}%）",size=10.2,bold=True,color=NAVY,space=2)
        title.paragraph_format.keep_with_next=True
        p(doc,f"日期：{m['sent_at']}　發報：{m['sender']}　收報：{m['recipient']}",size=8.3,space=2)
        if m['observations']:
            for o in m['observations']:
                p(doc,f"輔助資訊：{o['observed_at']}／{o['source']}／{o['activity']}／{o['location']}／可信度 {o['confidence']}／{o['status']}",size=8.1,color=GREY,space=2)
        else:p(doc,'輔助資訊：此封無第三方目擊；有收發者、日期與跨報共現線索。',size=8.1,color=GREY,space=2)
        p(doc,'原始密電：'+' '.join(codes),size=8.3,space=2)
        p(doc,'最終還原：'+'｜'.join('［未解］' if x=='[未解]' else x for x in guessed),size=8.2,space=2)
        p(doc,'正確明文：'+'｜'.join(words),size=8.2,space=3)
        mid=(len(codes)+1)//2
        def detail(j):
            if j>=len(codes):return '—'
            guessed_word='［未解］' if guessed[j]=='[未解]' else guessed[j]
            mark='✓' if hit[j] else '✗'
            return f'{j+1:02d}　{codes[j]}　{mark}　{guessed_word} → {words[j]}'
        table(doc,['位置／碼組／結果／猜測→真值（前半）','位置／碼組／結果／猜測→真值（後半）'],
             [[detail(i),detail(i+mid)] for i in range(mid)],fontsize=7.7)
        if index%2==0 and index!=len(msgs):doc.add_page_break()
    assert count_correct==470 and count_total==852
    out.parent.mkdir(parents=True,exist_ok=True)
    doc.save(out)
    return out

if __name__=='__main__':
    ap=argparse.ArgumentParser()
    ap.add_argument('--data-dir',type=Path,required=True)
    ap.add_argument('--out',type=Path,required=True)
    args=ap.parse_args()
    output=render(args.data_dir,args.out)
    print('Written:',output,'bytes:',output.stat().st_size)
